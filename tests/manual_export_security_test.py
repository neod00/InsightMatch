"""T11 F-04·F-05, T10 D-3(NFR-7) 회귀 — 매뉴얼 내보내기(PDF/DOCX).

- F-04: 마크다운 → ReportLab 변환에서 입력의 `<`, `&` 등을 이스케이프한다. `<img src=...>` 가 있어도
  로컬 파일·네트워크를 열지 않고 글자로 찍힌다. /api/export-iso 는 manual_id 가 있으면 저장본만 쓴다.
- F-05: 500 응답에 예외 원문을 넣지 않는다.
- NFR-7: 표지(첫 페이지) 하단에 서버 고정 고지문.

외부 호출·서버 기동 없음. 실제 PDF/DOCX 변환은 로컬 라이브러리(reportlab, python-docx)로만 한다.
"""
import builtins
import io
import os
import shutil
import tempfile
import unittest
from unittest.mock import MagicMock, patch

import docx
from reportlab.lib import utils as rl_utils
from reportlab.lib.pagesizes import A4
from reportlab.platypus import Paragraph, SimpleDocTemplate

from tests.manual_retry_test import ManualRetryTestBase, auth_headers

from index import db  # noqa: E402
from models import ManualGeneration  # noqa: E402
from services import document_export_service as des  # noqa: E402

try:
    import pypdf
except ImportError:  # 선택 의존성 — 없으면 PDF 텍스트 추출 검사만 건너뛴다
    pypdf = None

PRD_NFR7_TEXT = '본 문서는 ISO 표준 원문을 대체하지 않으며, 인증 심사 전 표준 원문과 대조 검토가 필요합니다'


class ResourceSpy:
    """변환 중 열린 파일 경로와 URL 읽기 시도를 기록한다 (실제 네트워크 요청은 보내지 않음)."""

    def __init__(self):
        self.opened = []
        self.urls = []
        self._real_open = builtins.open

    def _open(self, file, *args, **kwargs):
        if isinstance(file, (str, bytes, os.PathLike)):
            self.opened.append(os.fsdecode(file))
        return self._real_open(file, *args, **kwargs)

    def _url_read(self, name, *args, **kwargs):
        self.urls.append(name)
        raise OSError('network disabled in test')

    def __enter__(self):
        self._patches = [patch('builtins.open', self._open),
                         patch.object(rl_utils, 'rlUrlRead', self._url_read)]
        for p in self._patches:
            p.start()
        return self

    def __exit__(self, *exc):
        for p in reversed(self._patches):
            p.stop()

    def touched(self, needle):
        needle = os.path.normcase(needle)
        return [p for p in self.opened if needle in os.path.normcase(p)]


class ProbeMixin:
    def make_probe(self):
        self.tmpdir = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.tmpdir, ignore_errors=True)
        self.probe = os.path.join(self.tmpdir, 'probe_secret.ini').replace('\\', '/')
        with open(self.probe, 'w', encoding='utf-8') as fp:
            fp.write('[secret]\nkey=value\n')

    def injected_markdown(self):
        img_file = f'<img src="{self.probe}" width="10" height="10"/>'
        img_net = '<img src="http://169.254.169.254/latest/meta-data" width="10" height="10"/>'
        return (f'# 제목 {img_file}\n\n'
                f'본문 {img_file} & <a href="http://example.invalid/">링크</a> **굵게**\n\n'
                f'- 항목 {img_net}\n'
                f'1. 순번 {img_file}\n\n'
                f'| 헤더 {img_file} | b |\n|---|---|\n| <font face="{self.probe}">셀</font> | {img_net} |\n')


class RichInlineEscapeTest(unittest.TestCase):
    def test_markup_is_escaped_before_our_tags(self):
        out = des.rich_inline_markdown('<img src="C:/Windows/win.ini"/> & **굵게** *기울임* `코드`')
        self.assertTrue(out.startswith('&lt;img src="C:/Windows/win.ini"/&gt; &amp; '), out)
        self.assertIn('<b>굵게</b>', out)
        self.assertIn('<i>기울임</i>', out)
        self.assertIn('<font face="Courier" size="8">코드</font>', out)
        self.assertNotIn('<img', out)

    def test_paragraph_parse_has_no_image_fragment(self):
        """ReportLab 파싱 결과에 이미지 조각(cbDefn)이 없고, 태그는 글자로 남는다."""
        style = des._get_pdf_styles()['ISOBody']
        para = Paragraph(des.rich_inline_markdown('앞 <img src="C:/Windows/win.ini" width="9" height="9"/> 뒤'), style)
        self.assertFalse(any(getattr(f, 'cbDefn', None) for f in para.frags))
        text = ''.join(getattr(f, 'text', '') for f in para.frags)
        self.assertIn('<img src="C:/Windows/win.ini"', text)

    def test_broken_nesting_falls_back_to_plain_text(self):
        style = des._get_pdf_styles()['ISOBody']
        para = des._safe_paragraph('**a *b** c* <img src="x"/>', style)
        text = ''.join(getattr(f, 'text', '') for f in para.frags)
        self.assertIn('<img src="x"/>', text)


class PdfResourceLoadingTest(ProbeMixin, unittest.TestCase):
    def setUp(self):
        self.make_probe()

    def test_harness_detects_raw_markup_loading(self):
        """대조군: 이스케이프 없는 Paragraph 는 실제로 probe 파일을 연다 → 아래 검사가 의미 있음을 확인."""
        style = des._get_pdf_styles()['ISOBody']
        buf = io.BytesIO()
        with ResourceSpy() as spy:
            with self.assertRaises(Exception):
                SimpleDocTemplate(buf, pagesize=A4).build(
                    [Paragraph(f'x <img src="{self.probe}" width="10" height="10"/>', style)])
        self.assertTrue(spy.touched('probe_secret.ini'))

    def test_pdf_export_does_not_open_files_or_urls_from_markdown(self):
        with ResourceSpy() as spy:
            buf = des.markdown_to_pdf(self.injected_markdown(), '회사', 'ISO 9001:2015')
        self.assertTrue(buf.getvalue().startswith(b'%PDF'))
        self.assertEqual([], spy.touched('probe_secret.ini'))
        self.assertEqual([], spy.urls)

    @unittest.skipIf(pypdf is None, 'pypdf 미설치')
    def test_tags_are_rendered_as_text(self):
        buf = des.markdown_to_pdf(self.injected_markdown(), '회사', 'ISO 9001:2015')
        text = ''.join(page.extract_text() for page in pypdf.PdfReader(buf).pages)
        self.assertIn('<img src=', text)
        self.assertIn('<a href="http://example.invalid/">', text)

    def test_docx_export_does_not_open_files_or_urls(self):
        with ResourceSpy() as spy:
            buf = des.markdown_to_docx(self.injected_markdown(), '회사', 'ISO 9001:2015')
        self.assertEqual([], spy.touched('probe_secret.ini'))
        self.assertEqual([], spy.urls)
        body = '\n'.join(p.text for p in docx.Document(buf).paragraphs)
        self.assertIn('<img src=', body)


class CoverNoticeTest(unittest.TestCase):
    def test_notice_text_is_prd_nfr7(self):
        self.assertIn(PRD_NFR7_TEXT, des.COVER_NOTICE_TEXT)

    def test_pdf_first_page_draws_notice(self):
        canvas = MagicMock()
        des._first_page(canvas, MagicMock(page=1))
        drawn = [c.args[2] for c in canvas.drawCentredString.call_args_list]
        self.assertIn(des.COVER_NOTICE_TEXT, drawn)

    def test_pdf_build_uses_cover_callback_only_on_first_page(self):
        captured = {}

        def fake_build(self_doc, story, **kwargs):
            captured.update(kwargs)

        with patch.object(SimpleDocTemplate, 'build', fake_build):
            des.markdown_to_pdf('# 본문', '회사', 'ISO 9001:2015')
        self.assertIs(captured['onFirstPage'], des._first_page)
        self.assertIs(captured['onLaterPages'], des._add_page_number)

    @unittest.skipIf(pypdf is None, 'pypdf 미설치')
    def test_pdf_notice_on_first_page_only(self):
        md = '# 표지\n\n' + '\n\n'.join(f'문단 {i} ' + '내용 ' * 80 for i in range(40))
        reader = pypdf.PdfReader(des.markdown_to_pdf(md, '회사', 'ISO 9001:2015'))
        self.assertGreater(len(reader.pages), 1)
        first = reader.pages[0].extract_text().replace('\n', '')
        self.assertIn(PRD_NFR7_TEXT, first)
        self.assertNotIn(PRD_NFR7_TEXT, reader.pages[1].extract_text().replace('\n', ''))

    def test_docx_first_page_footer_has_notice(self):
        doc = docx.Document(des.markdown_to_docx('# 표지\n\n본문', '회사', 'ISO 9001:2015'))
        section = doc.sections[0]
        self.assertTrue(section.different_first_page_header_footer)
        first_footer = '\n'.join(p.text for p in section.first_page_footer.paragraphs)
        self.assertIn(des.COVER_NOTICE_TEXT, first_footer)
        self.assertNotIn(PRD_NFR7_TEXT, '\n'.join(p.text for p in section.footer.paragraphs))
        # 본문(LLM 생성)과 분리된 서버 템플릿이다
        self.assertNotIn(PRD_NFR7_TEXT, '\n'.join(p.text for p in doc.paragraphs))


class ExportRouteTest(ProbeMixin, ManualRetryTestBase):
    STORED = '# 저장본 제목\n\n저장된 본문'

    def export(self, body, user=None, fake=True):
        calls = {}

        def capture(name):
            def _fn(markdown, company_name='', target_iso=''):
                calls[name] = markdown
                return io.BytesIO(b'FILE')
            return _fn

        if fake:
            with patch('services.document_export_service.markdown_to_pdf', capture('pdf')), \
                    patch('services.document_export_service.markdown_to_docx', capture('docx')):
                r = self.client.post('/api/export-iso', json=body, headers=auth_headers(user or self.company))
        else:
            r = self.client.post('/api/export-iso', json=body, headers=auth_headers(user or self.company))
        return r, calls

    def stored_manual(self, phase1=STORED):
        session = self.new_session()
        m = self.manual(session)
        m.phase1_markdown = phase1
        db.session.commit()
        return session['manual_id']

    def test_manual_id_uses_stored_markdown_not_request_body(self):
        manual_id = self.stored_manual()
        r, calls = self.export({'manual_id': manual_id, 'format': 'pdf',
                                'markdown': '<img src="C:/Windows/win.ini"/> 요청 본문'})
        self.assertEqual(200, r.status_code)
        self.assertEqual(self.STORED, calls['pdf'])

    def test_request_markdown_used_only_without_manual_id(self):
        r, calls = self.export({'format': 'docx', 'markdown': '# 요청 본문'})
        self.assertEqual(200, r.status_code)
        self.assertEqual('# 요청 본문', calls['docx'])

    def test_manual_id_without_stored_markdown_is_400(self):
        session = self.new_session()
        r, calls = self.export({'manual_id': session['manual_id'], 'format': 'pdf', 'markdown': '# 요청 본문'})
        self.assertEqual(400, r.status_code)
        self.assertEqual({}, calls)

    def test_other_users_or_malformed_manual_id_is_404(self):
        manual_id = self.stored_manual()
        r, _ = self.export({'manual_id': manual_id, 'format': 'pdf'}, user=self.other)
        self.assertEqual(404, r.status_code)
        r, _ = self.export({'manual_id': {'$ne': None}, 'format': 'pdf', 'markdown': '# x'})
        self.assertEqual(404, r.status_code)

    def test_size_limit_still_applies_without_manual_id(self):
        r, _ = self.export({'format': 'pdf', 'markdown': 'a' * 250001})
        self.assertEqual(413, r.status_code)

    def test_500_does_not_leak_exception_text(self):
        def boom(*args, **kwargs):
            raise RuntimeError("cannot identify image file 'C:/secret/win.ini'")

        with patch('services.document_export_service.markdown_to_pdf', boom):
            r = self.client.post('/api/export-iso', json={'format': 'pdf', 'markdown': '# x'},
                                 headers=auth_headers(self.company))
        self.assertEqual(500, r.status_code)
        self.assertEqual('EXPORT_FAILED', r.get_json()['code'])
        text = r.get_data(as_text=True)
        self.assertNotIn('secret', text)
        self.assertNotIn('win.ini', text)

    def test_end_to_end_injection_without_manual_id_opens_nothing(self):
        self.make_probe()
        with ResourceSpy() as spy:
            r, _ = self.export({'format': 'pdf', 'markdown': self.injected_markdown()}, fake=False)
        self.assertEqual(200, r.status_code)
        self.assertTrue(r.data.startswith(b'%PDF'))
        self.assertEqual([], spy.touched('probe_secret.ini'))
        self.assertEqual([], spy.urls)

    def test_invalid_format_type_is_400(self):
        r, _ = self.export({'format': ['pdf'], 'markdown': '# x'})
        self.assertEqual(400, r.status_code)


if __name__ == '__main__':
    unittest.main()
